import io
from pathlib import Path

import pytest

import pii
from parsing import ParseError, parse_input

SAMPLE_CV = (Path(__file__).resolve().parent.parent / "samples" / "sample_cv.txt").read_text(encoding="utf-8")
ALL = (".pdf", ".doc", ".docx", ".txt")


def _docx_bytes(with_table=False) -> bytes:
    import docx

    d = docx.Document()
    for line in SAMPLE_CV.splitlines():
        d.add_paragraph(line)
    if with_table:
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Skill", "Excel"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in SAMPLE_CV.splitlines():
        c.drawString(40, y, line[:110])
        y -= 14
        if y < 40:
            c.showPage()
            y = 800
    c.save()
    return buf.getvalue()


def test_txt():
    doc = parse_input("cv.txt", SAMPLE_CV.encode("utf-8"), "", ALL)
    assert "Northwind" in doc.text and doc.source == "cv.txt"


def test_txt_cp1252():
    doc = parse_input("cv.txt", (SAMPLE_CV + " café").encode("cp1252"), "", ALL)
    assert "café" in doc.text


def test_docx_with_table_flag():
    doc = parse_input("cv.docx", _docx_bytes(with_table=True), "", ALL)
    assert "Demand Planning" in doc.text
    assert "Skill | Excel" in doc.text
    assert any("table" in f for f in doc.format_flags)


def test_pdf():
    doc = parse_input("cv.pdf", _pdf_bytes(), "", ALL)
    assert "Northwind" in doc.text


def test_file_preferred_over_paste():
    doc = parse_input("cv.txt", SAMPLE_CV.encode(), "pasted " * 100, ALL)
    assert doc.warnings and "Northwind" in doc.text


def test_pasted_text():
    assert parse_input(None, None, SAMPLE_CV, ALL).source == "pasted text"


def test_too_short_and_empty():
    with pytest.raises(ParseError, match="too short"):
        parse_input(None, None, "hello", ALL)
    with pytest.raises(ParseError, match="No input"):
        parse_input(None, None, "   ", ALL)


def test_disallowed_extension():
    with pytest.raises(ParseError, match="Unsupported"):
        parse_input("cv.txt", b"x" * 500, "", (".pdf", ".docx"))


def test_doc_without_libreoffice_gives_friendly_message(monkeypatch):
    import parsing

    monkeypatch.setattr(parsing, "_find_soffice", lambda: None)
    with pytest.raises(ParseError, match="DOCX or PDF"):
        parse_input("cv.doc", b"\xd0\xcf\x11\xe0" + b"0" * 600, "", ALL)


def test_corrupt_pdf():
    with pytest.raises(ParseError, match="Couldn't read"):
        parse_input("cv.pdf", b"not a pdf" * 100, "", ALL)


# ---------- PII ----------

def test_redaction():
    text, log = pii.redact(SAMPLE_CV, name="Ananya Rao", location="Mumbai, Maharashtra")
    assert "ananya.rao.sample@example.com" not in text and "[EMAIL]" in text
    assert "98765" not in text and "[PHONE]" in text
    assert "linkedin.com" not in text
    assert "Ananya" not in text and "[CANDIDATE]" in text
    assert "Date of Birth" not in text and "Marital" not in text
    # date ranges and other numbers survive
    assert "2021 - May 2025" in text and "38% to 27%" in text and "60,000 sq ft" in text
    assert any("protected" in x for x in log)
    assert pii.restore_name("Hi [CANDIDATE]", "Ananya Rao") == "Hi Ananya Rao"


def test_url_redaction_stops_at_quotes_and_brackets():
    text, _ = pii.redact('see "linkedin.com/in/x" and (https://a.example/b), [www.c.example]')
    assert text == 'see "[LINK]" and ([LINK]), [[LINK]]'


def test_redact_data_keeps_structure():
    data = {"contact": {"name": "Ananya Rao", "links": ["linkedin.com/in/ananya-rao"]},
            "skills": ["Excel"], "n": 3, "bullets": ['Built "dashboard" at https://x.example/y']}
    out = pii.redact_data(data, name="Ananya Rao", extra=["linkedin.com/in/ananya-rao"])
    assert out["contact"]["name"] == "[CANDIDATE]"
    assert "linkedin" not in str(out) and out["n"] == 3 and out["skills"] == ["Excel"]
    assert out["bullets"] == ['Built "dashboard" at [LINK]']
