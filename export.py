"""Report downloads: full Markdown, JSON, and the tailored CV as DOCX."""

import io
import json
import re
from typing import Dict, Optional

import pii
from config import FIT_DIMENSIONS, FIT_DISCLAIMER
from pipeline import Report, critical_gap_count


def report_json(rep: Report, scores: Dict[str, Optional[float]]) -> bytes:
    d = rep.to_dict()
    d["scores"] = scores
    return json.dumps(d, ensure_ascii=False, indent=2).encode("utf-8")


def _facts_md(title: str, res, dropped: int) -> list[str]:
    out = [f"## {title}"]
    if not res:
        return out + ["_Not available._", ""]
    out.append(f"_Subject: {res.subject}. Facts are cited; analysis and recommendations are AI-generated._\n")
    out += [f"- **[{f.topic}]** {f.statement} - [{f.source_title or 'source'}]({f.source_url})"
            + (f" ({f.published_date})" if f.published_date else "") for f in res.verified_facts]
    if dropped:
        out.append(f"\n_{dropped} item(s) removed because their source couldn't be verified._")
    out += ["", "**Analysis (AI)**"] + [f"- {a}" for a in res.analysis]
    out += ["", "**What this means for your interview (AI recommendation)**"] + [f"- {a}" for a in res.interview_implications]
    return out + [""]


def report_markdown(rep: Report, scores: Dict[str, Optional[float]]) -> bytes:
    jd = rep.jd
    L = [f"# CareerMatch AI report: {jd.job_title if jd else ''} at {jd.company if jd and jd.company else 'n/a'}",
         f"_Generated {rep.created} with {rep.model}. Career guidance only - not a hiring decision._", ""]

    def s(v):
        return "n/a" if v is None else f"{v:.0f}/100"

    L += ["## Scores",
          f"- Job fit: **{s(scores.get('fit'))}** - {FIT_DISCLAIMER}",
          f"- ATS readiness: **{s(scores.get('ats'))}**",
          f"- Interview readiness (heuristic): **{s(scores.get('readiness'))}**", ""]

    if rep.fit:
        L += ["## Fit by dimension", "| Dimension | Score | Rationale |", "|---|---|---|"]
        L += [f"| {FIT_DIMENSIONS[d.dimension][0]} | {d.score}/10 | {d.rationale} |" for d in rep.fit.dimensions]
        L += ["", "## Requirement-by-requirement", "| Requirement | Must-have | Status | CV evidence |", "|---|---|---|---|"]
        L += [f"| {r.requirement} | {'yes' if r.must_have else ''} | {r.status} | {r.cv_evidence} |" for r in rep.fit.requirements]
        L.append("")
    if rep.sg:
        L += ["## Top strengths"] + [f"{i}. **{x.title}** - \"{x.evidence_quote}\" ({x.jd_requirement})"
                                    for i, x in enumerate(rep.sg.strengths, 1)]
        L += ["", f"## Gaps ({critical_gap_count(rep)} critical)"]
        for g in rep.sg.gaps:
            L += [f"### [{g.severity}] {g.title}", g.why_it_matters,
                  f"- Short term: {g.close_short_term}", f"- Long term: {g.close_long_term}",
                  f"- In the interview: {g.interview_framing}"]
            L += [f"- Resource: {r.name} ({r.url_or_search})" for r in g.resources]
        L.append("")
    if rep.cv_opt:
        L += ["## CV improvements", "### Weak statements"]
        for w in rep.cv_opt.weak_statements:
            L += [f"- ~~{w.original}~~", f"  - Issue: {w.issue}", f"  - Rewrite: {w.rewrite}"]
        L += ["", "### Missing keywords"] + [f"- {k.keyword} ({k.recommendation.replace('_', ' ')}): {k.note}"
                                              for k in rep.cv_opt.missing_keywords]
        L.append("")
    if rep.questions:
        L += ["## Interview questions"]
        for q in sorted(rep.questions.questions, key=lambda q: ["high", "medium", "low"].index(q.priority)):
            L += [f"### [{q.priority}] {q.question}", f"_Category: {q.category}_",
                  f"- Evaluating: {q.evaluating}", f"- Answer guidance: {q.answer_guidance}", f"- Pitfalls: {q.pitfalls}"]
        L += ["", "### Self-feedback guidance"] + [f"- {x}" for x in rep.questions.feedback_guidance] + [""]
    if rep.cases:
        L += ["## Case interview practice"]
        for c in rep.cases.cases:
            L += [f"### {c.title} ({c.case_type.replace('_', ' ')})", c.prompt, "", "Clarifying questions:"]
            L += [f"- {x}" for x in c.clarifying_questions] + ["", "Framework:"]
            for b in c.framework:
                L += [f"- {b.branch}"] + [f"  - {p}" for p in b.sub_points]
            L += ["", "Model answer outline:"] + [f"{i}. {x}" for i, x in enumerate(c.model_answer_outline, 1)] + [""]
    L += _facts_md("Company research", rep.company, rep.company_dropped)
    L += _facts_md("Industry research", rep.industry, rep.industry_dropped)
    if rep.coaching:
        L += ["## Coaching answers"]
        for a in rep.coaching.answers:
            L += [f"### {a.question}", a.answer, f"_Tip: {a.tips}_", ""]
        L += ["## STAR story bank"]
        for st in rep.coaching.stories:
            L += [f"### {st.theme.replace('_', ' ').title()}: {st.title}",
                  f"_Source: {st.cv_source or 'not in CV - needs your input'}_",
                  f"- **S:** {st.situation}", f"- **T:** {st.task}", f"- **A:** {st.action}", f"- **R:** {st.result}", ""]
        L += ["## CV red flags"] + [f"- **{r.flag}**: {r.why_it_may_concern} -> {r.honest_response}"
                                    for r in rep.coaching.red_flags] + [""]
    if rep.plan:
        L += ["## 7-day preparation plan"]
        for d in rep.plan.days:
            L += [f"### Day {d.day}: {d.focus} ({d.time_estimate})"] + [f"- [ ] {t}" for t in d.tasks]
            L += [f"- Practice: {', '.join(d.practice)}", f"- Revision: {d.revision_focus}", ""]
        L += ["## Top 3 actions"] + [f"{i}. **{a.action}** - {a.rationale}" for i, a in enumerate(rep.plan.top_actions, 1)]
    if rep.cv_opt:
        L += ["", "## Tailored CV", "", pii.restore_name(rep.cv_opt.tailored_cv_markdown, rep.candidate_name)]
    return "\n".join(L).encode("utf-8")


def _add_runs(par, text: str):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part:
            par.add_run(part)


def tailored_cv_docx(markdown: str, name: str) -> bytes:
    """Minimal Markdown -> DOCX (headings, bullets, bold). Single column, ATS-friendly."""
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    for raw in pii.restore_name(markdown, name).splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", line):
            # '#' -> title (level 0), '##' -> Heading 1, '###' -> Heading 2
            d.add_heading(m.group(2).replace("**", ""), level=len(m.group(1)) - 1)
        elif m := re.match(r"^\s*[-*•]\s+(.*)", line):
            _add_runs(d.add_paragraph(style="List Bullet"), m.group(1))
        elif re.fullmatch(r"\s*[-_*]{3,}\s*", line):
            continue
        else:
            _add_runs(d.add_paragraph(), line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
