"""File -> text extraction for PDF, DOCX, DOC and TXT, plus ATS formatting-risk flags."""

import io
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from typing import List

MIN_CHARS = 300


class ParseError(Exception):
    """Raised with a user-facing message when a file can't be read."""


@dataclass
class ParsedDoc:
    text: str
    source: str  # "pasted text" or file name
    format_flags: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


# ---------- per-format readers ----------


def _read_pdf(data: bytes) -> tuple[str, List[str]]:
    flags: List[str] = []
    try:
        import pdfplumber

        pages, tables, images = [], 0, 0
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
                tables += len(page.find_tables())
                images += len(page.images)
        text = "\n".join(pages)
        if tables:
            flags.append(f"{tables} table(s) detected - many ATS parsers read tables poorly")
        if images:
            flags.append(f"{images} image(s)/graphic(s) detected - ATS ignores image content")
    except Exception:
        from pypdf import PdfReader

        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    return text, flags


def _read_docx(data: bytes) -> tuple[str, List[str]]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    parts: List[str] = []
    # Walk the body in order so table text stays next to its surrounding paragraphs.
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(Paragraph(child, d).text)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and t not in cells:  # merged cells repeat their text
                        cells.append(t)
                if cells:
                    parts.append(" | ".join(cells))
    flags: List[str] = []
    if d.tables:
        flags.append(f"{len(d.tables)} table(s) used for layout - many ATS parsers read tables poorly")
    header_text = " ".join(p.text for s in d.sections for p in s.header.paragraphs).strip()
    if header_text:
        flags.append("content in the page header - some ATS skip headers (keep contact details in the body)")
    if d.inline_shapes:
        flags.append(f"{len(d.inline_shapes)} image(s) detected - ATS ignores image content")
    return "\n".join(parts), flags


def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        if path := shutil.which(name):
            return path
    for path in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(path):
            return path
    return None


def _read_doc(data: bytes) -> tuple[str, List[str]]:
    soffice = _find_soffice()
    if not soffice:
        raise ParseError(
            "Old .doc files need LibreOffice to convert, and it isn't installed. "
            "Please save the file as DOCX or PDF, or paste the text instead."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "input.doc")
        with open(src, "wb") as f:
            f.write(data)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, src],
            check=False, capture_output=True, timeout=120,
        )
        out = os.path.join(tmp, "input.docx")
        if not os.path.exists(out):
            raise ParseError("Couldn't convert this .doc file. Please save it as DOCX or PDF, or paste the text.")
        with open(out, "rb") as f:
            return _read_docx(f.read())


def _read_txt(data: bytes) -> tuple[str, List[str]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc), []
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), []


READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".doc": _read_doc, ".txt": _read_txt}


# ---------- text-level checks ----------

_ODD_CHARS = set("★☆✓✔✗✘➤►▶◆◇■□●○❖⚫️🔹🔸")


def _text_flags(text: str) -> List[str]:
    flags = []
    odd = sum(text.count(c) for c in _ODD_CHARS)
    if odd > 15:
        flags.append(f"{odd} decorative symbols - may garble in some ATS")
    long_lines = sum(1 for l in text.splitlines() if len(l) > 250)
    if long_lines > 3:
        flags.append("several very long lines - possible multi-column layout merged during parsing")
    return flags


def parse_input(filename: str | None, data: bytes | None, pasted: str, allowed: tuple[str, ...]) -> ParsedDoc:
    """Prefer an uploaded file over pasted text. Raises ParseError with a friendly message."""
    warnings: List[str] = []
    if data:
        ext = os.path.splitext(filename or "")[1].lower()
        if ext not in allowed:
            raise ParseError(f"Unsupported file type '{ext}'. Allowed: {', '.join(allowed)}")
        if pasted.strip():
            warnings.append("Both a file and pasted text were given - using the uploaded file.")
        try:
            text, flags = READERS[ext](data)
        except ParseError:
            raise
        except Exception as e:  # corrupt file, password-protected PDF, etc.
            raise ParseError(f"Couldn't read '{filename}' ({type(e).__name__}). Try another format or paste the text.")
        source = filename or "uploaded file"
    else:
        text, flags, source = pasted, [], "pasted text"

    text = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").splitlines()).strip()
    if not text:
        raise ParseError("No input provided.")
    if len(text) < MIN_CHARS:
        hint = " This looks like a scanned/image PDF - paste the text instead." if source.lower().endswith(".pdf") else ""
        raise ParseError(f"Only {len(text)} characters were extracted from {source}; that's too short to analyse.{hint}")
    return ParsedDoc(text=text, source=source, format_flags=flags + _text_flags(text), warnings=warnings)
